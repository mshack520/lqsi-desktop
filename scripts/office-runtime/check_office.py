"""Open and inspect a supported Office file with the bundled Python libraries."""
import argparse
import json
from pathlib import Path


def docx_text(path):
    from docx import Document

    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


def pptx_text(path):
    from pptx import Presentation

    presentation = Presentation(path)
    chunks = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                chunks.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def xlsx_text(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return "\n".join(
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        )
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--contains")
    args = parser.parse_args()
    extension = args.path.suffix.lower()
    readers = {".docx": docx_text, ".pptx": pptx_text, ".xlsx": xlsx_text}
    if extension not in readers:
        raise ValueError(f"Unsupported Office file: {extension or '(no extension)'}")
    text = readers[extension](args.path)
    passed = args.contains is None or args.contains in text
    print(json.dumps({"verdict": "pass" if passed else "fail", "format": extension[1:]}, ensure_ascii=False))
    if not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
